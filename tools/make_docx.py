"""Render a project Markdown doc into a font-compliant .docx.

Fonts (per AGENTS.md, for Chinese documents):
    headings   -> 黑体 (SimHei)
    body text  -> 微软雅黑 (Microsoft YaHei)
    code/JSON  -> Consolas   (allowed for structured/code text)

    python tools/make_docx.py                       # DESIGN.md -> docs/音乐解密工具-设计文档.docx
    python tools/make_docx.py SRC.md OUT.docx       # explicit paths

Supports: #/##/### headings, ``` code fences, bullet/numbered lists,
> blockquotes, --- rules, pipe tables, inline `code` and **bold**.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SRC = ROOT / "DESIGN.md"
DEFAULT_OUT = ROOT / "docs" / "音乐解密工具-设计文档.docx"

BODY_FONT = "微软雅黑"
HEAD_FONT = "黑体"
MONO_FONT = "Consolas"

_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_TABLE_SEP = re.compile(r"^\s*\|[\s:\-|]+\|\s*$")


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None,
                 color: tuple | None = None) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_runs(p, text: str, *, font: str, size: float, bold: bool = False, color=None) -> None:
    """Add `text` to paragraph `p`, honouring inline `code` and **bold**."""
    for piece in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = p.add_run(piece[1:-1])
            set_run_font(r, MONO_FONT, size - 1, False, (0xB0, 0x30, 0x30))
        elif piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            r = p.add_run(piece[2:-2])
            set_run_font(r, font, size, True, color)
        else:
            r = p.add_run(piece)
            set_run_font(r, font, size, bold, color)


def add_para(doc, text: str, *, font=BODY_FONT, size=10.5, bold=False, color=None,
             space_after=6, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.3
    if indent is not None:
        p.paragraph_format.left_indent = Pt(indent)
    add_runs(p, text, font=font, size=size, bold=bold, color=color)
    return p


def add_code_block(doc, lines) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.left_indent = Pt(10)
        r = p.add_run(line if line else " ")
        set_run_font(r, MONO_FONT, 8.5, False, (0x22, 0x22, 0x22))


def add_table(doc, rows: "list[list[str]]") -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            p = cell.paragraphs[0]
            for r in list(p.runs):                 # drop any empty default run
                r._element.getparent().remove(r._element)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, row[j] if j < len(row) else "",
                     font=BODY_FONT, size=9.5, bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def convert(src: Path, out: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    lines = src.read_text("utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):                      # code fence
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            add_code_block(doc, block)
            i += 1
            continue

        if _TABLE_ROW.match(line):                      # pipe table
            rows = []
            while i < len(lines) and _TABLE_ROW.match(lines[i]):
                if not _TABLE_SEP.match(lines[i]):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue

        if line.startswith("### "):
            add_para(doc, line[4:], font=HEAD_FONT, size=12, bold=True, space_after=4)
        elif line.startswith("## "):
            add_para(doc, line[3:], font=HEAD_FONT, size=14, bold=True, space_after=6)
        elif line.startswith("# "):
            add_para(doc, line[2:], font=HEAD_FONT, size=18, bold=True, space_after=10,
                     color=(0x1F, 0x3A, 0x93))
        elif line.startswith("> "):
            add_para(doc, line[2:], size=9.5, color=(0x55, 0x55, 0x55), indent=14)
        elif re.match(r"^\s*[-*] ", line):
            add_para(doc, "• " + re.sub(r"^\s*[-*] ", "", line), indent=16, space_after=2)
        elif re.match(r"^\s*\d+\. ", line):
            add_para(doc, re.sub(r"^\s*", "", line), indent=16, space_after=2)
        elif line.strip() == "---":
            add_para(doc, "─" * 46, size=8, color=(0xBB, 0xBB, 0xBB), space_after=8)
        elif line.strip() == "":
            pass
        else:
            add_para(doc, line)
        i += 1

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"wrote {out}")


def main(argv: "list[str]") -> int:
    src = Path(argv[1]) if len(argv) > 1 else DEFAULT_SRC
    out = Path(argv[2]) if len(argv) > 2 else DEFAULT_OUT
    convert(src, out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))